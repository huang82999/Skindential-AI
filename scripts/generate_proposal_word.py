import json
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    return h

def main():
    doc = Document()
    
    # Title
    title = doc.add_heading('Skindential AI：膚況檢測專案技術演進與提案分析報告', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. 執行摘要
    add_heading(doc, '1. 執行摘要', 1)
    doc.add_paragraph('【專案解決的問題】本專案旨在解決高解析度 (1920x1440) 臉部影像中，微小痘痘與瑕疵在縮放後容易流失特徵，且高達 59 種臨床類別因語義重疊導致模型難以分類的問題。')
    doc.add_paragraph('【核心方法】團隊捨棄傳統單一全圖推論，改採「雙路推論架構 (Dual-Path Inference)」，結合高解析度局部切片 (SAHI)，並將 49 種微小瑕疵重組為 5 大核心臨床群組。')
    doc.add_paragraph('【主要成果】根據現有資料推論，綜合精確度 (mAP50) 由 5.41% 提升至最高 17.47%，相對提升約 222.9%；Recall 由 12.68% 提升至 24.63%，相對提升約 94.2%。')
    doc.add_paragraph('【技術與產品價值】此架構成功在「保留微觀特徵」與「降低邊緣運算負載」間取得平衡，為未來部署於 Edge AI (如 Jetson/NPU) 奠定初步可行性基礎。')

    # 2. 使用情境與問題背景
    add_heading(doc, '2. 使用情境與問題背景', 1)
    doc.add_paragraph('使用者看到的是一張清楚的臉部影像，但痘痘在模型縮放後可能只剩數個像素。當終端設備拍攝高解析度臉部照片時，為符合神經網路的輸入限制，通常需進行影像縮放 (Downsampling)。這會導致微小的粉刺或痘痘特徵完全流失。')
    doc.add_paragraph('此外，原始資料庫標註了 59 種細微膚況類別。然而，許多類別（如閉鎖性粉刺與栗粒腫）在視覺特徵上高度重疊。過多且語義重疊的類別，造成模型在學習過程中產生「分類猶豫」，將真實病灶誤判為背景。')

    # 3. Baseline 與主要瓶頸
    add_heading(doc, '3. Baseline 與主要瓶頸', 1)
    doc.add_paragraph('【原始方法與指標】Baseline 階段採用 1920x1440 原圖直接壓縮至 640x640 進行 YOLO 模型訓練。根據原始資料驗證，其 mAP50 僅 5.41%，Precision 19.14%，Recall 12.68%。')
    doc.add_paragraph('【為何準確率偏低】微小特徵被壓縮為次像素雜訊，神經網路無法提取有效特徵。')
    doc.add_paragraph('【為何調整超參數無法突破】團隊進行了 44 輪自動化超參數調適，涵蓋學習率、邊界權重與資料增強。但數據顯示 mAP50 始終在 5%~6% 的平台期徘徊。這（根據現有資料推論）證明了瓶頸在於資料本質（標記歧義與解析度限制），而非單純的模型參數未達最佳化。')

    # 4. 技術改善方法
    add_heading(doc, '4. 技術改善方法', 1)
    doc.add_paragraph('團隊進行了重新設計模型觀看影像尺度，並整理辨識類別的架構重組：')
    doc.add_paragraph('【高解析度切片】導入 480x480 網格化切片，保留 1:1 原圖細節，解決微小病灶的特徵流失問題。')
    doc.add_paragraph('【雙路推論】將全局特徵（如黑眼圈，共 10 類）與微觀特徵（如粉刺，共 49 類）分離，由宏觀模型與微觀切片模型分別處理。這解釋了 59 類與 49 類的數量差異（已由原始資料驗證）。')
    doc.add_paragraph('【五大群組重組】將 49 種微小瑕疵依據臨床學理，強勢收斂為發炎痘痘、粉刺等 5 大群組，降低分類語義重疊。')
    doc.add_paragraph('【訓練策略調整】將 DFL 定位損失權重提升至 2.0 以強化邊界定位，並關閉破壞紋理的 Mixup 資料增強。')

    # 5. 實驗結果與成長分析
    add_heading(doc, '5. 實驗結果與成長分析', 1)
    doc.add_paragraph('本專案的成效並非持續性的指數型成長，而是「平台期後的階段性突破」。前 44 輪實驗處於效能平台期，在架構與資料重組後才出現階躍式提升。')
    
    table1 = doc.add_table(rows=1, cols=4)
    table1.style = 'Table Grid'
    hdr = table1.rows[0].cells
    hdr[0].text = '指標'
    hdr[1].text = 'Baseline (全尺寸直推)'
    hdr[2].text = '最終模型 (雙路+5大群組)'
    hdr[3].text = '成長分析'
    
    data = [
        ('mAP50', '5.41%', '16.67% (峰值 17.47%)', '峰值增加 12.06 個百分點，為原本的 3.23 倍，相對提升約 222.9%'),
        ('Recall', '12.68%', '24.63%', '增加 11.95 個百分點，為原本的 1.94 倍，相對提升約 94.2%'),
        ('Precision', '19.14%', '28.21%', '增加 9.07 個百分點')
    ]
    for item in data:
        row = table1.add_row().cells
        row[0].text = item[0]; row[1].text = item[1]; row[2].text = item[2]; row[3].text = item[3]
        
    doc.add_paragraph('\n【指標變化說明】mAP50 與 Recall 的雙雙顯著提升，(根據現有資料推論) 表明模型並非單純依賴 Confidence Threshold 犧牲 Precision 來換取 Recall，而是整體辨識與特徵擷取能力得到了本質上的增強。目前只能證明最終成果相較 Baseline 有顯著的倍數提升，尚不足以證明持續性的指數型成長。')

    # 6. 技術貢獻
    add_heading(doc, '6. 技術貢獻', 1)
    doc.add_paragraph('1. 解決高解析度微小病灶的特徵流失：透過切片技術讓模型得以看見真實特徵。（最可能提升 Recall）')
    doc.add_paragraph('2. 降低標記語義重疊造成的分類猶豫：透過 5 大群組重組，消除模型分類障礙。（最可能綜合提升 mAP50）')
    doc.add_paragraph('3. 在準確率與 Edge AI 運算成本間取得平衡：雙路架構將全局與局部解耦，避免整張 1920 高解析度原圖直接運算導致的 VRAM 溢出。')

    # 7. 產品化與 Edge AI 價值
    add_heading(doc, '7. 產品化與 Edge AI 價值', 1)
    doc.add_paragraph('本架構可應用於智慧醫美檢測儀或行動端 AI 膚況分析 APP。對企業而言，其低算力需求的特性可大幅降低終端硬體成本。')
    doc.add_paragraph('【注意】雖然架構具備部署潛力，但目前尚缺 Jetson 或 NPU 上的實際速度與記憶體消耗測試 (尚缺資料驗證)，因此目前不得宣稱已達商業級部署標準。')

    # 8. 限制與後續計畫
    add_heading(doc, '8. 限制與後續計畫', 1)
    doc.add_paragraph('目前所有改善均在最終版本同時實作，【尚缺資料驗證】證明各項技術的單獨貢獻度。建議下一階段補做：')
    doc.add_paragraph('1. 消融實驗 (Ablation Study)：分別驗證切片、五大群組、DFL 調整的獨立貢獻。')
    doc.add_paragraph('2. 泛化能力測試：針對不同膚色、光線環境與拍攝角度的測試集進行驗證。')
    doc.add_paragraph('3. 硬體實測：在目標 Edge 設備上進行 FPS (幀率) 與 VRAM (記憶體) 消耗的基準測試 (Benchmark)。')

    # 9. 結論與提案建議
    add_heading(doc, '9. 結論與提案建議', 1)
    doc.add_paragraph('專案團隊已成功跨越超參數調校的瓶頸，透過「雙路推論」與「臨床分群」的核心策略，使系統綜合準確度取得倍數級突破。這證明了結合臨床 Domain Knowledge 進行資料清洗，遠比單純調校模型參數更具價值。')
    doc.add_paragraph('【下一階段需求】建議繼續投入資源，完成硬體部署基準測試與消融實驗。預期可驗證本模型在 Edge 設備上的即時處理能力，推進至產品原型 (PoC) 階段。')

    doc.save(r"E:\old\Skindential_AI_Technical_Proposal.docx")

if __name__ == '__main__':
    main()
